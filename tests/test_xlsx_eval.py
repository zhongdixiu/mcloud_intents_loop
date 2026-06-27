from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from intent_router.skills import SkillRegistry
from intent_router.types import DialogueHistory, RouteResult, SkillRef
from intent_router.xlsx_eval import (
    HistoryCaseTurn,
    build_gold_dialogue_history,
    default_output_path,
    evaluate_format_xlsx_cases,
    evaluate_xlsx_cases,
    load_format_xlsx_cases,
    load_xlsx_cases,
    normalize_code,
    result_matches_expected_codes,
    split_codes,
)
from intent_router.xlsx_eval import _build_code_index


def test_normalize_and_split_codes() -> None:
    assert normalize_code(0) == "000"
    assert normalize_code("0000") == "000"
    assert normalize_code(999) == "018"
    assert normalize_code("999") == "018"
    assert normalize_code(1) == "036006"
    assert normalize_code("001") == "036006"
    assert normalize_code(28) == "040001"
    assert normalize_code("028") == "040001"
    assert normalize_code(32) == "040003"
    assert normalize_code("032") == "040003"
    assert normalize_code(22) == "022"
    assert normalize_code("22") == "022"
    assert normalize_code(36006) == "036006"
    assert split_codes("012、22, 036006/0/999/001/028/032") == [
        "012",
        "022",
        "036006",
        "000",
        "018",
        "040001",
        "040003",
    ]


def test_result_matches_022_only_for_mcloud_search_01x() -> None:
    search_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="搜图片",
        code="012",
    )
    note_result = RouteResult(
        status="matched",
        skill=SkillRef(id="note_skill", name="AI笔记"),
        intent="搜笔记",
        code="017",
    )

    assert result_matches_expected_codes(
        search_result,
        expected_code="022",
        alternate_codes=[],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": True, "reason": "022_search_equivalent"}
    assert result_matches_expected_codes(
        note_result,
        expected_code="022",
        alternate_codes=[],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": False, "reason": "code_mismatch"}


def test_result_matching_ignores_022_equivalence_when_expected_has_many_codes() -> None:
    search_012_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="搜图片",
        code="012",
    )
    search_014_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="搜视频",
        code="014",
    )
    search_016_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="搜音频",
        code="016",
    )
    search_023_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="历史意图023",
        code="023",
    )

    assert result_matches_expected_codes(
        search_012_result,
        expected_code="022",
        alternate_codes=["014", "016", "023"],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": False, "reason": "code_mismatch"}
    assert result_matches_expected_codes(
        search_014_result,
        expected_code="022",
        alternate_codes=["014", "016", "023"],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": True, "reason": "exact_code"}
    assert result_matches_expected_codes(
        search_016_result,
        expected_code="022",
        alternate_codes=["014", "016", "023"],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": False, "reason": "code_mismatch"}
    assert result_matches_expected_codes(
        search_023_result,
        expected_code="022",
        alternate_codes=["014", "016", "023"],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": False, "reason": "code_mismatch"}


def test_result_matching_accepts_any_expected_code_when_many_codes_without_022() -> None:
    search_016_result = RouteResult(
        status="matched",
        skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
        intent="搜音频",
        code="016",
    )

    assert result_matches_expected_codes(
        search_016_result,
        expected_code="014",
        alternate_codes=["016", "023"],
        search_equivalent_codes={"012", "013", "014", "015", "016", "017", "018"},
    ) == {"matched": True, "reason": "exact_code"}


def test_load_xlsx_cases_skips_empty_history_cells(tmp_path: Path) -> None:
    path = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "第一轮对话",
            "第一轮预期意图",
            "第二轮对话",
            "第二轮预期意图",
            "当前对话",
            "当前预期意图",
            "备注意图",
        ],
    )
    sheet.append(["搜合同", "018", None, None, "蓝色", "022", "012、013"])
    workbook.save(path)

    cases = load_xlsx_cases(path)

    assert len(cases) == 1
    assert cases[0].row_index == 2
    assert cases[0].query == "蓝色"
    assert cases[0].expected_code == "012"
    assert cases[0].alternate_codes == []
    assert cases[0].source_expected_codes == ["022", "012", "013"]
    assert cases[0].history_turns == [
        HistoryCaseTurn(query="搜合同", expected_code="018"),
    ]


def test_load_format_xlsx_cases_uses_history_min_code_and_expected_codes(
    tmp_path: Path,
) -> None:
    path = tmp_path / "format_cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "上文1",
            "上文意图1",
            "上文2",
            "上文意图2",
            "对话",
            "预期意图",
        ],
    )
    sheet.append(
        [
            "搜索三国演义的小说",
            "022，013，016，023",
            "中国机长",
            "999",
            "权力的游戏",
            "999",
        ],
    )
    sheet.append(
        [
            None,
            None,
            None,
            None,
            "无间道的呐",
            "022，014，016，023",
        ],
    )
    workbook.save(path)

    cases = load_format_xlsx_cases(path)

    assert len(cases) == 2
    assert cases[0].query == "权力的游戏"
    assert cases[0].expected_code == "018"
    assert cases[0].alternate_codes == []
    assert cases[0].history_turns == [
        HistoryCaseTurn(query="搜索三国演义的小说", expected_code="013"),
        HistoryCaseTurn(query="中国机长", expected_code="018"),
    ]
    assert cases[1].expected_code == "014"
    assert cases[1].alternate_codes == []
    assert cases[1].source_expected_codes == ["022", "014", "016", "023"]


def test_build_gold_history_maps_022_to_search_context() -> None:
    registry = SkillRegistry.from_path("skills")
    history = build_gold_dialogue_history(
        [HistoryCaseTurn(query="搜试卷", expected_code="022")],
        registry,
        _build_code_index(registry),
    )

    turn = history.turns[0]
    assert turn.user_query == "搜试卷"
    assert turn.result.status == "matched"
    assert turn.result.skill_id == "mcloud_search_skill"
    assert turn.result.code == "022"


async def test_evaluate_xlsx_cases_writes_rows_and_summary(tmp_path: Path) -> None:
    path = tmp_path / "cases.xlsx"
    output = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "第一轮对话",
            "第一轮预期意图",
            "当前对话",
            "当前预期意图",
            "备注意图",
        ],
    )
    sheet.append(["搜合同文件", "018", "蓝色", "022", None])
    sheet.append([None, None, "闲聊一下", "000", None])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜图片",
                code="012",
                loop_count=2,
                correction_scopes=["intent_mismatch"],
            ),
            RouteResult(
                status="matched",
                skill=None,
                intent="普通对话",
                code="000",
                loop_count=1,
            ),
        ],
    )

    summary = await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
    )

    assert summary["total"] == 2
    assert summary["passed"] == 2
    assert summary["accuracy"] == 1.0
    assert summary["output_path"] == str(output)
    assert [len(history.turns) for history in router.histories] == [1, 0]

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    summary_sheet = workbook["summary"]
    headers = [cell.value for cell in result_sheet[1]]
    first = dict(zip(headers, [cell.value for cell in result_sheet[2]], strict=True))
    second = dict(zip(headers, [cell.value for cell in result_sheet[3]], strict=True))
    assert first["matched"] is True
    assert first["eval_mode"] == "gold_history"
    assert first["match_reason"] == "022_search_equivalent"
    assert first["loop_count"] == 2
    assert first["correction_scopes"] == '["intent_mismatch"]'
    assert second["predicted_code"] == "000"
    summary_rows = {
        row[0].value: row[1].value
        for row in summary_sheet.iter_rows(min_row=2, max_col=2)
    }
    assert summary_rows["passed"] == 2


async def test_evaluate_format_xlsx_cases_writes_results_with_new_matching_rule(
    tmp_path: Path,
) -> None:
    path = tmp_path / "format_cases.xlsx"
    output = tmp_path / "format_result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["上文1", "上文意图1", "对话", "期望意图标签"])
    sheet.append(["搜电影", "022", "无间道", "022，014，016，023"])
    sheet.append([None, None, "扫毒", "022"])
    sheet.append([None, None, "权力的游戏", "999"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜图片",
                code="012",
                loop_count=1,
            ),
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜图片",
                code="012",
                loop_count=1,
            ),
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜综合",
                code="018",
                loop_count=2,
            ),
        ],
    )

    summary = await evaluate_format_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
    )

    assert summary["total"] == 3
    assert summary["passed"] == 2
    assert summary["failed"] == 1
    assert [len(history.turns) for history in router.histories] == [1, 0, 0]

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    headers = [cell.value for cell in result_sheet[1]]
    rows = [
        dict(zip(headers, [cell.value for cell in row], strict=True))
        for row in result_sheet.iter_rows(min_row=2, max_row=4)
    ]
    assert rows[0]["matched"] is False
    assert rows[0]["match_reason"] == "code_mismatch"
    assert rows[0]["expected_code"] == "014"
    assert rows[0]["expected_codes"] == '["022", "014", "016", "023"]'
    assert rows[0]["effective_expected_codes"] == '["014"]'
    assert rows[1]["matched"] is True
    assert rows[1]["match_reason"] == "022_search_equivalent"
    assert rows[2]["matched"] is True
    assert rows[2]["expected_code"] == "018"


async def test_evaluate_xlsx_cases_defaults_output_to_same_directory(
    tmp_path: Path,
) -> None:
    path = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["当前对话", "当前预期意图"])
    sheet.append(["闲聊一下", "000"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=None,
                intent="普通对话",
                code="000",
                loop_count=1,
            ),
        ],
    )

    summary = await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        router=router,  # type: ignore[arg-type]
    )

    expected_output = default_output_path(path)
    assert summary["output_path"] == str(expected_output)
    assert expected_output.exists()


async def test_evaluate_xlsx_cases_reports_progress(tmp_path: Path) -> None:
    path = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["当前对话", "当前预期意图"])
    sheet.append(["闲聊一下", "000"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=None,
                intent="普通对话",
                code="000",
                loop_count=1,
            ),
        ],
    )
    progress_events = []

    await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        router=router,  # type: ignore[arg-type]
        progress_callback=lambda processed, total, record, passed: progress_events.append(
            (processed, total, record["matched"], passed),
        ),
    )

    assert progress_events == [(1, 1, True, 1)]


async def test_evaluate_xlsx_cases_records_row_exception_and_continues(
    tmp_path: Path,
) -> None:
    path = tmp_path / "cases.xlsx"
    output = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["当前对话", "当前预期意图"])
    sheet.append(["会触发异常", "000"])
    sheet.append(["闲聊一下", "000"])
    workbook.save(path)

    router = FakeRouter(
        [
            RuntimeError("structured output validation failed"),
            RouteResult(
                status="matched",
                intent="普通对话",
                code="000",
                loop_count=1,
            ),
        ],
    )

    summary = await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
    )

    assert summary["total"] == 2
    assert summary["passed"] == 1
    assert summary["failed"] == 1
    assert summary["error_count"] == 1

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    headers = [cell.value for cell in result_sheet[1]]
    rows = [
        dict(zip(headers, [cell.value for cell in row], strict=True))
        for row in result_sheet.iter_rows(min_row=2, max_row=3)
    ]
    assert rows[0]["predicted_status"] == "error"
    assert rows[0]["matched"] is False
    assert rows[0]["match_reason"] == "exception"
    assert rows[0]["error_type"] == "RuntimeError"
    assert rows[0]["error"] == "structured output validation failed"
    assert rows[1]["predicted_status"] == "matched"
    assert rows[1]["matched"] is True


async def test_evaluate_xlsx_cases_end_to_end_routes_history_before_current(
    tmp_path: Path,
) -> None:
    path = tmp_path / "cases.xlsx"
    output = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "第一轮对话",
            "第一轮预期意图",
            "第二轮对话",
            "第二轮预期意图",
            "当前对话",
            "当前预期意图",
        ],
    )
    sheet.append(["推荐刘德华的歌曲", "000", "有没有电影", "000", "找图片", "022"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(status="matched", intent="普通对话", code="000", loop_count=1),
            RouteResult(status="matched", intent="普通对话", code="000", loop_count=2),
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜图片",
                code="012",
                loop_count=1,
            ),
        ],
    )

    summary = await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
        if_end2end=True,
    )

    assert summary["total"] == 1
    assert summary["passed"] == 1
    assert summary["accuracy"] == 1.0
    assert summary["eval_mode"] == "end2end"
    assert summary["intent_only"] is False
    assert router.queries == ["推荐刘德华的歌曲", "有没有电影", "找图片"]
    assert [len(history.turns) for history in router.histories] == [0, 1, 2]

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    headers = [cell.value for cell in result_sheet[1]]
    row = dict(zip(headers, [cell.value for cell in result_sheet[2]], strict=True))
    assert row["eval_mode"] == "end2end"
    assert row["intent_only"] is False
    assert row["history_predicted_codes"] == '["000", "000"]'
    assert row["history_loop_counts"] == "[1, 2]"
    assert row["matched"] is True
    assert row["match_reason"] == "022_search_equivalent"
    turn_results = json.loads(row["turn_results"])
    assert [turn["query"] for turn in turn_results] == [
        "推荐刘德华的歌曲",
        "有没有电影",
        "找图片",
    ]
    assert [turn["is_scored"] for turn in turn_results] == [False, False, True]
    assert turn_results[-1]["predicted_code"] == "012"


async def test_evaluate_format_xlsx_cases_end_to_end_records_each_turn(
    tmp_path: Path,
) -> None:
    path = tmp_path / "format_cases.xlsx"
    output = tmp_path / "format_result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["上文1", "上文意图1", "对话", "期望意图标签"])
    sheet.append(["搜索谢娜的图片", "022", "再生成一些近期的", "036011"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜图片",
                code="012",
                loop_count=1,
            ),
            RouteResult(
                status="matched",
                skill=SkillRef(id="image_skill", name="图片处理"),
                intent="生成图片",
                code="036011",
                loop_count=2,
            ),
        ],
    )

    summary = await evaluate_format_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
        if_end2end=True,
    )

    assert summary["passed"] == 1
    assert router.queries == ["搜索谢娜的图片", "再生成一些近期的"]
    assert [len(history.turns) for history in router.histories] == [0, 1]

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    headers = [cell.value for cell in result_sheet[1]]
    row = dict(zip(headers, [cell.value for cell in result_sheet[2]], strict=True))
    assert row["predicted_code"] == "036011"
    turn_results = json.loads(row["turn_results"])
    assert len(turn_results) == 2
    assert turn_results[0]["is_scored"] is False
    assert turn_results[1]["is_scored"] is True


async def test_evaluate_xlsx_cases_records_intent_only_mode(tmp_path: Path) -> None:
    path = tmp_path / "cases.xlsx"
    output = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["当前对话", "当前预期意图"])
    sheet.append(["找相关文档", "013"])
    workbook.save(path)

    router = FakeRouter(
        [
            RouteResult(
                status="matched",
                skill=SkillRef(id="mcloud_search_skill", name="云盘搜索"),
                intent="搜文档",
                code="013",
                loop_count=1,
            ),
        ],
    )

    summary = await evaluate_xlsx_cases(
        path,
        skills_path="skills",
        output_path=output,
        router=router,  # type: ignore[arg-type]
        intent_only=True,
    )

    assert summary["intent_only"] is True

    workbook = load_workbook(output)
    result_sheet = workbook["results"]
    headers = [cell.value for cell in result_sheet[1]]
    row = dict(zip(headers, [cell.value for cell in result_sheet[2]], strict=True))
    assert row["intent_only"] is True


class FakeRouter:
    def __init__(self, results: list[RouteResult | Exception]) -> None:
        self.results = results
        self.histories: list[DialogueHistory] = []
        self.queries: list[str] = []

    async def route(
        self,
        query: str,
        *,
        dialogue_history: DialogueHistory | None = None,
        context: dict[str, Any] | None = None,
        trace: list[dict[str, Any]] | None = None,
    ) -> RouteResult:
        self.queries.append(query)
        self.histories.append(
            (dialogue_history or DialogueHistory()).model_copy(deep=True),
        )
        if trace is not None:
            trace.append({"event": "fake"})
        result = self.results.pop(0)
        if isinstance(result, Exception):
            raise result
        return result
