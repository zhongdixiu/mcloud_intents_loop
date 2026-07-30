from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from .config import RouterConfig
from .dialogue import IntentDialogueAgent
from .router import IntentRouter
from .skills import SkillRegistry
from .types import (
    DialogueHistory,
    DialogueRouteSummary,
    DialogueTurn,
    RouteResult,
    SkillDefinition,
    SkillRef,
)


CURRENT_QUERY_HEADER = "当前对话"
CURRENT_EXPECTED_HEADER = "当前预期意图"
CURRENT_ALTERNATE_HEADER = "备注意图"
FORMAT_CURRENT_QUERY_HEADER = "对话"
FORMAT_EXPECTED_HEADERS = ("期望意图标签", "预期意图")
EXPECTED_SKILL_HEADERS = ("expected_skill_id", "预期Skill ID")
EXPECTED_INTENT_HEADERS = ("expected_intent", "预期Intent")
EXPECTED_CODE_HEADERS = ("expected_code", "预期Code")
EXPECTED_PARAMS_HEADERS = ("expected_params", "预期参数")
ALTERNATE_ROUTES_HEADERS = ("alternate_routes", "备选Route")
SEARCH_SKILL_ID = "mcloud_search_skill"
SEARCH_022_CODE = "022"
GENERIC_MULTI_ACCEPT_CODES = {"016", "022"}
ORDINARY_DIALOGUE_CODE = "000"
LEGACY_999_EQUIVALENT_CODE = "018"
LEGACY_CODE_EQUIVALENTS = {
    "001": "036006",
    "028": "040001",
    "032": "040003",
    "999": LEGACY_999_EQUIVALENT_CODE,
}
STRONG_BUSINESS_ACTIONS = (
    "搜索",
    "搜",
    "查找",
    "找",
    "打开",
    "进入",
    "入口",
    "工具",
    "发送",
    "整理",
    "筛选",
    "生成",
    "创作",
    "编辑",
    "处理",
    "翻译",
    "总结",
    "识别",
)
MULTI_INTENT_RESOURCE_TERMS = (
    "文档",
    "文件",
    "图片",
    "照片",
    "视频",
    "电影",
    "音频",
    "音乐",
    "歌曲",
    "文件夹",
)
MULTI_INTENT_JOINERS = ("和", "与", "及", "以及", "还有", "、", "，", ",", "/")


@dataclass(frozen=True)
class HistoryCaseTurn:
    query: str
    expected_code: str


@dataclass(frozen=True)
class XlsxEvalCase:
    row_index: int
    query: str
    expected_code: str
    alternate_codes: list[str]
    history_turns: list[HistoryCaseTurn]
    source_expected_codes: list[str] = field(default_factory=list)
    allow_022_search_equivalence: bool = False
    expected_skill_id: str | None = None
    expected_intent: str | None = None
    expected_params: dict[str, Any] | None = None
    alternate_routes: list[dict[str, str]] = field(default_factory=list)


async def evaluate_xlsx_cases(
    cases_path: Path,
    *,
    skills_path: str | Path = "skills",
    output_path: Path | None = None,
    trace_enabled: bool = False,
    router: IntentRouter | None = None,
    dialogue_history_limit: int | None = 5,
    if_end2end: bool = False,
    compare_no_loop: bool = False,
    evaluator_mode: str = "conditional",
    progress_callback: Callable[[int, int, dict[str, Any], int], None] | None = None,
) -> dict[str, Any]:
    registry = SkillRegistry.from_path(skills_path)
    if router is None:
        config = RouterConfig.model_validate(
            {"decision": {"evaluator_mode": evaluator_mode}},
        )
        router = IntentRouter.from_config(
            skills_path=skills_path,
            dialogue_history_limit=dialogue_history_limit,
            config=config,
        )
    cases = load_xlsx_cases(cases_path)
    output_path = normalize_output_path(output_path or default_output_path(cases_path))
    return await _evaluate_loaded_cases(
        cases,
        registry=registry,
        router=router,
        output_path=output_path,
        trace_enabled=trace_enabled,
        if_end2end=if_end2end,
        compare_no_loop=compare_no_loop,
        progress_callback=progress_callback,
    )


async def evaluate_format_xlsx_cases(
    cases_path: Path,
    *,
    skills_path: str | Path = "skills",
    output_path: Path | None = None,
    trace_enabled: bool = False,
    router: IntentRouter | None = None,
    dialogue_history_limit: int | None = 5,
    if_end2end: bool = False,
    compare_no_loop: bool = False,
    evaluator_mode: str = "conditional",
    progress_callback: Callable[[int, int, dict[str, Any], int], None] | None = None,
) -> dict[str, Any]:
    registry = SkillRegistry.from_path(skills_path)
    if router is None:
        config = RouterConfig.model_validate(
            {"decision": {"evaluator_mode": evaluator_mode}},
        )
        router = IntentRouter.from_config(
            skills_path=skills_path,
            dialogue_history_limit=dialogue_history_limit,
            config=config,
        )
    cases = load_format_xlsx_cases(cases_path)
    output_path = normalize_output_path(output_path or default_output_path(cases_path))
    return await _evaluate_loaded_cases(
        cases,
        registry=registry,
        router=router,
        output_path=output_path,
        trace_enabled=trace_enabled,
        if_end2end=if_end2end,
        compare_no_loop=compare_no_loop,
        progress_callback=progress_callback,
    )


async def _evaluate_loaded_cases(
    cases: list[XlsxEvalCase],
    *,
    registry: SkillRegistry,
    router: IntentRouter,
    output_path: Path,
    trace_enabled: bool,
    if_end2end: bool,
    compare_no_loop: bool,
    progress_callback: Callable[[int, int, dict[str, Any], int], None] | None,
) -> dict[str, Any]:
    code_index = _build_code_index(registry)
    search_equivalent_codes = _search_022_equivalent_codes(registry)

    total = 0
    passed = 0
    elapsed_values: list[float] = []
    total_elapsed_values: list[float] = []
    loop_counts: list[int] = []
    status_counts: dict[str, int] = {}
    no_loop_passed = 0
    no_loop_elapsed_values: list[float] = []
    no_loop_total_elapsed_values: list[float] = []
    no_loop_status_counts: dict[str, int] = {}
    both_passed = 0
    both_failed = 0
    loop_only_passed = 0
    no_loop_only_passed = 0
    records: list[dict[str, Any]] = []
    first_candidate_passed = 0
    top_n_recall_passed = 0
    full_route_passed = 0
    evaluator_switch_count = 0
    evaluator_switch_gain = 0
    evaluator_switch_loss = 0
    expansion_count = 0
    expansion_gain = 0
    expansion_waste = 0
    evaluator_call_count = 0
    unnecessary_evaluator_call_count = 0
    model_call_counts: list[int] = []
    retry_count = 0
    timeout_count = 0
    parameter_labeled_count = 0
    parameter_exact_match_count = 0
    skill_recall_at_1 = 0
    skill_recall_at_3 = 0
    skill_recall_at_5 = 0
    intent_recall_at_1 = 0
    intent_recall_at_2 = 0
    ordinary_dialogue_count = 0
    business_action_to_000_count = 0
    strict_passed = 0
    source_expected_passed = 0
    multi_code_case_count = 0
    multi_code_adjusted_count = 0

    for case in cases:
        total += 1
        multi_code_case_count += int(_is_multi_code_case(case))
        multi_code_adjusted_count += int(_is_multi_code_adjusted(case))
        case_started = time.perf_counter()
        try:
            evaluation = await _run_eval_case(
                case,
                registry=registry,
                code_index=code_index,
                router=router,
                trace_enabled=trace_enabled,
                if_end2end=if_end2end,
                no_loop=False,
            )
        except Exception as exc:
            elapsed_ms = (time.perf_counter() - case_started) * 1000
            elapsed_values.append(elapsed_ms)
            total_elapsed_values.append(elapsed_ms)
            loop_counts.append(0)
            status_counts["error"] = status_counts.get("error", 0) + 1
            record = _build_error_record(
                case=case,
                error=exc,
                elapsed_ms=elapsed_ms,
                total_elapsed_ms=elapsed_ms,
                eval_mode="end2end" if if_end2end else "gold_history",
            )
            if compare_no_loop:
                no_loop_started = time.perf_counter()
                try:
                    no_loop_evaluation = await _run_eval_case(
                        case,
                        registry=registry,
                        code_index=code_index,
                        router=router,
                        trace_enabled=trace_enabled,
                        if_end2end=if_end2end,
                        no_loop=True,
                    )
                except Exception as no_loop_exc:
                    no_loop_elapsed_ms = (
                        time.perf_counter() - no_loop_started
                    ) * 1000
                    no_loop_elapsed_values.append(no_loop_elapsed_ms)
                    no_loop_total_elapsed_values.append(no_loop_elapsed_ms)
                    no_loop_status_counts["error"] = (
                        no_loop_status_counts.get("error", 0) + 1
                    )
                    no_loop_matched = False
                    record.update(
                        _build_no_loop_error_fields(
                            error=no_loop_exc,
                            elapsed_ms=no_loop_elapsed_ms,
                            total_elapsed_ms=no_loop_elapsed_ms,
                        ),
                    )
                else:
                    no_loop_result = no_loop_evaluation["result"]
                    no_loop_elapsed_ms = no_loop_evaluation["elapsed_ms"]
                    no_loop_total_elapsed_ms = no_loop_evaluation["total_elapsed_ms"]
                    no_loop_elapsed_values.append(no_loop_elapsed_ms)
                    no_loop_total_elapsed_values.append(no_loop_total_elapsed_ms)
                    no_loop_status_counts[no_loop_result.status] = (
                        no_loop_status_counts.get(no_loop_result.status, 0) + 1
                    )
                    no_loop_match = result_matches_expected_codes(
                        no_loop_result,
                        expected_code=case.expected_code,
                        alternate_codes=_case_alternate_codes(case),
                        search_equivalent_codes=search_equivalent_codes,
                        allow_022_search_equivalence=case.allow_022_search_equivalence,
                    )
                    no_loop_matched = bool(no_loop_match["matched"])
                    no_loop_passed += int(no_loop_matched)
                    record.update(
                        _build_no_loop_record_fields(
                            result=no_loop_result,
                            matched=no_loop_matched,
                            match_reason=no_loop_match["reason"],
                            elapsed_ms=no_loop_elapsed_ms,
                            total_elapsed_ms=no_loop_total_elapsed_ms,
                            turn_results=no_loop_evaluation["turn_results"],
                            trace=(
                                no_loop_evaluation["trace"] if trace_enabled else None
                            ),
                            trace_summary=no_loop_evaluation["trace_summary"],
                        ),
                    )
                both_failed += int(not no_loop_matched)
                no_loop_only_passed += int(no_loop_matched)
            records.append(record)
            if progress_callback is not None:
                progress_callback(total, len(cases), record, passed)
            continue

        result = evaluation["result"]
        elapsed_ms = evaluation["elapsed_ms"]
        total_elapsed_ms = evaluation["total_elapsed_ms"]
        elapsed_values.append(elapsed_ms)
        total_elapsed_values.append(total_elapsed_ms)
        loop_counts.append(result.loop_count)
        status_counts[result.status] = status_counts.get(result.status, 0) + 1

        match = result_matches_expected_codes(
            result,
            expected_code=case.expected_code,
            alternate_codes=_case_alternate_codes(case),
            search_equivalent_codes=search_equivalent_codes,
            allow_022_search_equivalence=case.allow_022_search_equivalence,
        )
        loop_matched = bool(match["matched"])
        passed += int(loop_matched)
        strict_match = _result_matches_case_codes(
            result,
            case=case,
            expected_codes=strict_expected_codes_for_matching(
                _source_expected_codes(case),
            ),
            search_equivalent_codes=search_equivalent_codes,
        )
        source_expected_match = _result_matches_case_codes(
            result,
            case=case,
            expected_codes=source_expected_codes_for_matching(
                _source_expected_codes(case),
            ),
            search_equivalent_codes=search_equivalent_codes,
        )
        strict_passed += int(strict_match["matched"])
        source_expected_passed += int(source_expected_match["matched"])
        candidate_metrics = _candidate_metrics(
            result,
            case=case,
            code_index=code_index,
            search_equivalent_codes=search_equivalent_codes,
        )
        first_candidate_passed += int(candidate_metrics["first_candidate_matched"])
        top_n_recall_passed += int(candidate_metrics["top_n_recalled"])
        full_route_passed += int(candidate_metrics["full_route_matched"])
        skill_recall_at_1 += int(candidate_metrics["skill_recall_at_1"])
        skill_recall_at_3 += int(candidate_metrics["skill_recall_at_3"])
        skill_recall_at_5 += int(candidate_metrics["skill_recall_at_5"])
        intent_recall_at_1 += int(candidate_metrics["intent_recall_at_1"])
        intent_recall_at_2 += int(candidate_metrics["intent_recall_at_2"])
        evaluator_switch_count += int(candidate_metrics["evaluator_switched"])
        evaluator_switch_gain += int(candidate_metrics["evaluator_switch_gain"])
        evaluator_switch_loss += int(candidate_metrics["evaluator_switch_loss"])
        expansion_count += int(candidate_metrics["expanded"])
        expansion_gain += int(candidate_metrics["expansion_gain"])
        expansion_waste += int(candidate_metrics["expansion_waste"])
        evaluator_call_count += int(candidate_metrics["evaluator_called"])
        unnecessary_evaluator_call_count += int(
            candidate_metrics["unnecessary_evaluator_call"],
        )
        model_calls = (
            result.diagnostics.model_calls
            if result.diagnostics is not None
            else []
        )
        model_call_counts.append(len(model_calls))
        retry_count += sum(max(0, call.attempts - 1) for call in model_calls)
        timeout_count += sum(call.status == "timeout" for call in model_calls)
        if case.expected_params is not None:
            parameter_labeled_count += 1
            parameter_exact_match_count += int(result.params == case.expected_params)
        ordinary_dialogue_count += int(result.code == ORDINARY_DIALOGUE_CODE)
        business_action_to_000_count += int(
            result.code == ORDINARY_DIALOGUE_CODE and _has_strong_business_action(case.query),
        )

        record = _build_record(
            case=case,
            result=result,
            matched=loop_matched,
            match_reason=match["reason"],
            elapsed_ms=elapsed_ms,
            total_elapsed_ms=total_elapsed_ms,
            eval_mode="end2end" if if_end2end else "gold_history",
            turn_results=evaluation["turn_results"],
            trace_summary=evaluation["trace_summary"],
            candidate_metrics=candidate_metrics,
            strict_matched=bool(strict_match["matched"]),
            strict_match_reason=strict_match["reason"],
            source_expected_matched=bool(source_expected_match["matched"]),
            source_expected_match_reason=source_expected_match["reason"],
        )
        if trace_enabled:
            record["trace"] = evaluation["trace"]
        if compare_no_loop:
            no_loop_started = time.perf_counter()
            try:
                no_loop_evaluation = await _run_eval_case(
                    case,
                    registry=registry,
                    code_index=code_index,
                    router=router,
                    trace_enabled=trace_enabled,
                    if_end2end=if_end2end,
                    no_loop=True,
                )
            except Exception as exc:
                no_loop_elapsed_ms = (time.perf_counter() - no_loop_started) * 1000
                no_loop_elapsed_values.append(no_loop_elapsed_ms)
                no_loop_total_elapsed_values.append(no_loop_elapsed_ms)
                no_loop_status_counts["error"] = (
                    no_loop_status_counts.get("error", 0) + 1
                )
                no_loop_matched = False
                record.update(
                    _build_no_loop_error_fields(
                        error=exc,
                        elapsed_ms=no_loop_elapsed_ms,
                        total_elapsed_ms=no_loop_elapsed_ms,
                    ),
                )
            else:
                no_loop_result = no_loop_evaluation["result"]
                no_loop_elapsed_ms = no_loop_evaluation["elapsed_ms"]
                no_loop_total_elapsed_ms = no_loop_evaluation["total_elapsed_ms"]
                no_loop_elapsed_values.append(no_loop_elapsed_ms)
                no_loop_total_elapsed_values.append(no_loop_total_elapsed_ms)
                no_loop_status_counts[no_loop_result.status] = (
                    no_loop_status_counts.get(no_loop_result.status, 0) + 1
                )
                no_loop_match = result_matches_expected_codes(
                    no_loop_result,
                    expected_code=case.expected_code,
                    alternate_codes=_case_alternate_codes(case),
                    search_equivalent_codes=search_equivalent_codes,
                    allow_022_search_equivalence=case.allow_022_search_equivalence,
                )
                no_loop_matched = bool(no_loop_match["matched"])
                no_loop_passed += int(no_loop_matched)
                record.update(
                    _build_no_loop_record_fields(
                        result=no_loop_result,
                        matched=no_loop_matched,
                        match_reason=no_loop_match["reason"],
                        elapsed_ms=no_loop_elapsed_ms,
                        total_elapsed_ms=no_loop_total_elapsed_ms,
                        turn_results=no_loop_evaluation["turn_results"],
                        trace=no_loop_evaluation["trace"] if trace_enabled else None,
                        trace_summary=no_loop_evaluation["trace_summary"],
                    ),
                )

            both_passed += int(loop_matched and no_loop_matched)
            both_failed += int(not loop_matched and not no_loop_matched)
            loop_only_passed += int(loop_matched and not no_loop_matched)
            no_loop_only_passed += int(not loop_matched and no_loop_matched)
        records.append(record)
        if progress_callback is not None:
            progress_callback(total, len(cases), record, passed)

    failed = total - passed
    summary = {
        "total": total,
        "passed": passed,
        "failed": failed,
        "accuracy": (passed / total) if total else 0.0,
        "adjusted_code_accuracy": (passed / total) if total else 0.0,
        "strict_accuracy": (strict_passed / total) if total else 0.0,
        "source_expected_accuracy": (
            source_expected_passed / total
        ) if total else 0.0,
        "eval_mode": "end2end" if if_end2end else "gold_history",
        "avg_elapsed_ms": _average(elapsed_values),
        "p50_elapsed_ms": _percentile(elapsed_values, 50),
        "p95_elapsed_ms": _percentile(elapsed_values, 95),
        "p99_elapsed_ms": _percentile(elapsed_values, 99),
        "avg_total_elapsed_ms": _average(total_elapsed_values),
        "avg_loop_count": _average(loop_counts),
        "matched_count": status_counts.get("matched", 0),
        "clarify_count": status_counts.get("clarify", 0),
        "abstain_count": status_counts.get("abstain", 0),
        "unsupported_count": status_counts.get("unsupported", 0),
        "error_count": status_counts.get("error", 0),
        "output_path": str(output_path),
        "first_candidate_passed": first_candidate_passed,
        "first_candidate_accuracy": (first_candidate_passed / total) if total else 0.0,
        "final_accuracy": (passed / total) if total else 0.0,
        "top_n_recall": (top_n_recall_passed / total) if total else 0.0,
        "full_route_passed": full_route_passed,
        "full_route_accuracy": (full_route_passed / total) if total else 0.0,
        "skill_recall_at_1": skill_recall_at_1 / total if total else 0.0,
        "skill_recall_at_3": skill_recall_at_3 / total if total else 0.0,
        "skill_recall_at_5": skill_recall_at_5 / total if total else 0.0,
        "intent_recall_at_1": intent_recall_at_1 / total if total else 0.0,
        "intent_recall_at_2": intent_recall_at_2 / total if total else 0.0,
        "evaluator_switch_count": evaluator_switch_count,
        "evaluator_switch_gain": evaluator_switch_gain,
        "evaluator_switch_loss": evaluator_switch_loss,
        "evaluator_switch_net_gain": evaluator_switch_gain - evaluator_switch_loss,
        "wrong_switch_rate": (
            evaluator_switch_loss / evaluator_switch_count
            if evaluator_switch_count
            else 0.0
        ),
        "evaluator_call_count": evaluator_call_count,
        "evaluator_call_rate": (
            evaluator_call_count / total if total else 0.0
        ),
        "unnecessary_evaluator_call_count": unnecessary_evaluator_call_count,
        "unnecessary_evaluator_call_rate": (
            unnecessary_evaluator_call_count / evaluator_call_count
            if evaluator_call_count
            else 0.0
        ),
        "expansion_count": expansion_count,
        "expansion_gain": expansion_gain,
        "expansion_waste": expansion_waste,
        "expansion_waste_rate": (
            expansion_waste / expansion_count if expansion_count else 0.0
        ),
        "avg_model_calls": _average(model_call_counts),
        "retry_count": retry_count,
        "timeout_count": timeout_count,
        "parameter_labeled_count": parameter_labeled_count,
        "parameter_exact_match_count": parameter_exact_match_count,
        "parameter_exact_match": (
            parameter_exact_match_count / parameter_labeled_count
            if parameter_labeled_count
            else None
        ),
        "ordinary_dialogue_count": ordinary_dialogue_count,
        "business_action_to_000_count": business_action_to_000_count,
        "multi_code_case_count": multi_code_case_count,
        "multi_code_adjusted_count": multi_code_adjusted_count,
    }
    if compare_no_loop:
        no_loop_failed = total - no_loop_passed
        summary.update(
            {
                "compare_no_loop": True,
                "no_loop_passed": no_loop_passed,
                "no_loop_failed": no_loop_failed,
                "no_loop_accuracy": (no_loop_passed / total) if total else 0.0,
                "no_loop_avg_elapsed_ms": _average(no_loop_elapsed_values),
                "no_loop_avg_total_elapsed_ms": _average(
                    no_loop_total_elapsed_values,
                ),
                "no_loop_matched_count": no_loop_status_counts.get("matched", 0),
                "no_loop_clarify_count": no_loop_status_counts.get("clarify", 0),
                "no_loop_abstain_count": no_loop_status_counts.get("abstain", 0),
                "no_loop_unsupported_count": no_loop_status_counts.get(
                    "unsupported",
                    0,
                ),
                "no_loop_error_count": no_loop_status_counts.get("error", 0),
                "both_passed": both_passed,
                "both_failed": both_failed,
                "loop_only_passed": loop_only_passed,
                "no_loop_only_passed": no_loop_only_passed,
            },
        )
    write_xlsx_result(output_path, records, summary)
    return summary


async def _run_eval_case(
    case: XlsxEvalCase,
    *,
    registry: SkillRegistry,
    code_index: dict[str, tuple[SkillDefinition, str]],
    router: IntentRouter,
    trace_enabled: bool,
    if_end2end: bool,
    no_loop: bool,
) -> dict[str, Any]:
    if if_end2end:
        return await _run_end_to_end_case(
            case,
            router=router,
            trace_enabled=trace_enabled,
            no_loop=no_loop,
        )

    history = build_gold_dialogue_history(
        case.history_turns,
        registry,
        code_index,
    )
    return await _run_gold_history_case(
        case,
        router=router,
        history=history,
        trace_enabled=trace_enabled,
        no_loop=no_loop,
    )


async def _run_gold_history_case(
    case: XlsxEvalCase,
    *,
    router: IntentRouter,
    history: DialogueHistory,
    trace_enabled: bool,
    no_loop: bool = False,
) -> dict[str, Any]:
    trace: list[dict[str, Any]] = []
    started = time.perf_counter()
    route_func = router.route_no_loop if no_loop else router.route
    result = await route_func(case.query, dialogue_history=history, trace=trace)
    elapsed_ms = (time.perf_counter() - started) * 1000
    turn_results = [
        _build_turn_result(
            turn_index=len(case.history_turns) + 1,
            query=case.query,
            expected_code=case.expected_code,
            result=result,
            elapsed_ms=elapsed_ms,
            is_scored=True,
        ),
    ]
    return {
        "result": result,
        "elapsed_ms": elapsed_ms,
        "total_elapsed_ms": elapsed_ms,
        "turn_results": turn_results,
        "trace": trace if trace_enabled else None,
        "trace_summary": _trace_summary(trace),
    }


async def _run_end_to_end_case(
    case: XlsxEvalCase,
    *,
    router: IntentRouter,
    trace_enabled: bool,
    no_loop: bool = False,
) -> dict[str, Any]:
    agent = IntentDialogueAgent(router)
    turn_results: list[dict[str, Any]] = []
    turn_traces: list[dict[str, Any]] = []
    total_started = time.perf_counter()

    turns = [
        *[
            {
                "query": history_turn.query,
                "expected_code": history_turn.expected_code,
                "is_scored": False,
            }
            for history_turn in case.history_turns
        ],
        {
            "query": case.query,
            "expected_code": case.expected_code,
            "is_scored": True,
        },
    ]

    result: RouteResult | None = None
    final_elapsed_ms = 0.0
    for index, turn in enumerate(turns, start=1):
        trace: list[dict[str, Any]] = []
        started = time.perf_counter()
        send_func = agent.send_no_loop if no_loop else agent.send
        result = await send_func(turn["query"], trace=trace)
        elapsed_ms = (time.perf_counter() - started) * 1000
        if turn["is_scored"]:
            final_elapsed_ms = elapsed_ms
        turn_results.append(
            _build_turn_result(
                turn_index=index,
                query=turn["query"],
                expected_code=turn["expected_code"],
                result=result,
                elapsed_ms=elapsed_ms,
                is_scored=turn["is_scored"],
            ),
        )
        if trace_enabled:
            turn_traces.append(
                {
                    "turn_index": index,
                    "query": turn["query"],
                    "trace": trace,
                },
            )

    if result is None:
        raise RuntimeError("Excel 评测用例缺少当前对话，无法执行端到端评测")

    total_elapsed_ms = (time.perf_counter() - total_started) * 1000
    return {
        "result": result,
        "elapsed_ms": final_elapsed_ms,
        "total_elapsed_ms": total_elapsed_ms,
        "turn_results": turn_results,
        "trace": turn_traces if trace_enabled else None,
        "trace_summary": _trace_summary(trace or []),
    }


def default_output_path(cases_path: Path) -> Path:
    return cases_path.with_name(f"{cases_path.stem}_测试结果.xlsx")


def normalize_output_path(output_path: Path) -> Path:
    if not output_path.suffix:
        return output_path.with_suffix(".xlsx")
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError(f"结果文件必须是 .xlsx: {output_path}")
    return output_path


def load_xlsx_cases(cases_path: Path) -> list[XlsxEvalCase]:
    workbook = load_workbook(cases_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [_cell_text(cell) for cell in header_row]
    header_to_index = {
        header: index
        for index, header in enumerate(headers)
        if header
    }
    current_query_index = _require_header(header_to_index, CURRENT_QUERY_HEADER)
    current_expected_index = _require_any_header(
        header_to_index,
        (CURRENT_EXPECTED_HEADER, *EXPECTED_CODE_HEADERS),
    )
    current_alternate_index = header_to_index.get(CURRENT_ALTERNATE_HEADER)
    expected_skill_index = _first_header_index(header_to_index, EXPECTED_SKILL_HEADERS)
    expected_intent_index = _first_header_index(header_to_index, EXPECTED_INTENT_HEADERS)
    expected_params_index = _first_header_index(header_to_index, EXPECTED_PARAMS_HEADERS)
    alternate_routes_index = _first_header_index(
        header_to_index,
        ALTERNATE_ROUTES_HEADERS,
    )
    history_pairs = _history_column_pairs(headers, header_to_index)

    cases: list[XlsxEvalCase] = []
    for row_index, row in enumerate(rows, start=2):
        query = _cell_text(_row_value(row, current_query_index))
        expected_codes = split_codes(_row_value(row, current_expected_index))
        if not query or not expected_codes:
            continue

        if current_alternate_index is not None:
            expected_codes.extend(split_codes(_row_value(row, current_alternate_index)))
        expected_codes = _dedupe_codes(expected_codes)
        effective_expected_codes = expected_codes_for_matching(expected_codes)
        if not effective_expected_codes:
            continue

        history_turns = []
        for query_index, expected_index in history_pairs:
            history_query = _cell_text(_row_value(row, query_index))
            history_code = history_code_for_codes(
                split_codes(_row_value(row, expected_index)),
            )
            if not history_query or not history_code:
                continue
            history_turns.append(
                HistoryCaseTurn(query=history_query, expected_code=history_code),
            )

        cases.append(
            XlsxEvalCase(
                row_index=row_index,
                query=query,
                expected_code=effective_expected_codes[0],
                alternate_codes=effective_expected_codes[1:],
                history_turns=history_turns,
                source_expected_codes=expected_codes,
                allow_022_search_equivalence=allows_022_search_equivalence(
                    expected_codes,
                ),
                expected_skill_id=_optional_cell(row, expected_skill_index),
                expected_intent=_optional_cell(row, expected_intent_index),
                expected_params=_parse_expected_params(
                    _row_value(row, expected_params_index)
                    if expected_params_index is not None
                    else None,
                ),
                alternate_routes=_parse_alternate_routes(
                    _row_value(row, alternate_routes_index)
                    if alternate_routes_index is not None
                    else None,
                ),
            ),
        )
    return cases


def load_format_xlsx_cases(cases_path: Path) -> list[XlsxEvalCase]:
    workbook = load_workbook(cases_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [_cell_text(cell) for cell in header_row]
    header_to_index = {
        header: index
        for index, header in enumerate(headers)
        if header
    }
    current_query_index = _require_header(header_to_index, FORMAT_CURRENT_QUERY_HEADER)
    current_expected_index = _require_any_header(
        header_to_index,
        (*FORMAT_EXPECTED_HEADERS, *EXPECTED_CODE_HEADERS),
    )
    expected_skill_index = _first_header_index(header_to_index, EXPECTED_SKILL_HEADERS)
    expected_intent_index = _first_header_index(header_to_index, EXPECTED_INTENT_HEADERS)
    expected_params_index = _first_header_index(header_to_index, EXPECTED_PARAMS_HEADERS)
    alternate_routes_index = _first_header_index(
        header_to_index,
        ALTERNATE_ROUTES_HEADERS,
    )
    history_pairs = _format_history_column_pairs(headers, header_to_index)

    cases: list[XlsxEvalCase] = []
    for row_index, row in enumerate(rows, start=2):
        query = _cell_text(_row_value(row, current_query_index))
        expected_codes = split_codes(_row_value(row, current_expected_index))
        if not query or not expected_codes:
            continue
        expected_codes = _dedupe_codes(expected_codes)
        effective_expected_codes = expected_codes_for_matching(expected_codes)
        if not effective_expected_codes:
            continue

        history_turns = []
        for query_index, expected_index in history_pairs:
            history_query = _cell_text(_row_value(row, query_index))
            history_codes = split_codes(_row_value(row, expected_index))
            history_code = history_code_for_codes(history_codes)
            if not history_query or not history_code:
                continue
            history_turns.append(
                HistoryCaseTurn(query=history_query, expected_code=history_code),
            )

        cases.append(
            XlsxEvalCase(
                row_index=row_index,
                query=query,
                expected_code=effective_expected_codes[0],
                alternate_codes=effective_expected_codes[1:],
                history_turns=history_turns,
                source_expected_codes=expected_codes,
                allow_022_search_equivalence=allows_022_search_equivalence(
                    expected_codes,
                ),
                expected_skill_id=_optional_cell(row, expected_skill_index),
                expected_intent=_optional_cell(row, expected_intent_index),
                expected_params=_parse_expected_params(
                    _row_value(row, expected_params_index)
                    if expected_params_index is not None
                    else None,
                ),
                alternate_routes=_parse_alternate_routes(
                    _row_value(row, alternate_routes_index)
                    if alternate_routes_index is not None
                    else None,
                ),
            ),
        )
    return cases


def build_gold_dialogue_history(
    history_turns: list[HistoryCaseTurn],
    registry: SkillRegistry,
    code_index: dict[str, tuple[SkillDefinition, str]],
) -> DialogueHistory:
    turns: list[DialogueTurn] = []
    for history_turn in history_turns:
        summary = _summary_for_expected_code(
            history_turn.expected_code,
            registry,
            code_index,
        )
        turns.append(
            DialogueTurn(
                user_query=history_turn.query,
                result=summary,
                metadata={"source": "xlsx_gold_history"},
            ),
        )
    return DialogueHistory(turns=turns)


def result_matches_expected_codes(
    result: RouteResult,
    *,
    expected_code: str,
    alternate_codes: list[str],
    search_equivalent_codes: set[str],
    allow_022_search_equivalence: bool | None = None,
) -> dict[str, Any]:
    raw_accepted_codes = _dedupe_codes([expected_code, *alternate_codes])
    accepted_codes = expected_codes_for_matching(raw_accepted_codes)
    allow_search_equivalence = (
        allows_022_search_equivalence(raw_accepted_codes)
        if allow_022_search_equivalence is None
        else allow_022_search_equivalence
    )
    return _result_matches_codes(
        result,
        accepted_codes=accepted_codes,
        search_equivalent_codes=search_equivalent_codes,
        allow_022_search_equivalence=allow_search_equivalence,
    )


def _result_matches_case_codes(
    result: RouteResult,
    *,
    case: XlsxEvalCase,
    expected_codes: list[str],
    search_equivalent_codes: set[str],
) -> dict[str, Any]:
    return _result_matches_codes(
        result,
        accepted_codes=expected_codes,
        search_equivalent_codes=search_equivalent_codes,
        allow_022_search_equivalence=allows_022_search_equivalence(
            _source_expected_codes(case),
        ),
    )


def _result_matches_codes(
    result: RouteResult,
    *,
    accepted_codes: list[str],
    search_equivalent_codes: set[str],
    allow_022_search_equivalence: bool,
) -> dict[str, Any]:
    predicted_code = result.code
    predicted_skill_id = result.skill.id if result.skill else None

    for accepted_code in accepted_codes:
        if predicted_code == accepted_code:
            return {"matched": True, "reason": "exact_code"}
    if (
        allow_022_search_equivalence
        and predicted_skill_id == SEARCH_SKILL_ID
        and predicted_code in search_equivalent_codes
    ):
        return {"matched": True, "reason": "022_search_equivalent"}

    return {"matched": False, "reason": "code_mismatch"}


def split_codes(value: Any) -> list[str]:
    text = _cell_text(value)
    if not text:
        return []
    codes = []
    for item in re.split(r"[、,，;；/]", text):
        code = normalize_code(item)
        if code and code not in codes:
            codes.append(code)
    return codes


def history_code_for_codes(codes: list[str]) -> str | None:
    return min_code(_dedupe_codes(codes))


def expected_codes_for_matching(codes: list[str]) -> list[str]:
    deduped_codes = _dedupe_codes(codes)
    if len(deduped_codes) > 1:
        adjusted_codes = [
            code for code in deduped_codes if code not in GENERIC_MULTI_ACCEPT_CODES
        ]
        if adjusted_codes:
            return adjusted_codes
    return deduped_codes


def strict_expected_codes_for_matching(codes: list[str]) -> list[str]:
    deduped_codes = _dedupe_codes(codes)
    if len(deduped_codes) > 1 and SEARCH_022_CODE in deduped_codes:
        code = min_code(deduped_codes)
        return [code] if code else []
    return deduped_codes


def source_expected_codes_for_matching(codes: list[str]) -> list[str]:
    return _dedupe_codes(codes)


def allows_022_search_equivalence(codes: list[str]) -> bool:
    return _dedupe_codes(codes) == [SEARCH_022_CODE]


def _source_expected_codes(case: XlsxEvalCase) -> list[str]:
    return case.source_expected_codes or [case.expected_code, *case.alternate_codes]


def _case_alternate_codes(case: XlsxEvalCase) -> list[str]:
    return _dedupe_codes(
        [
            *case.alternate_codes,
            *[
                normalize_code(route.get("code"))
                for route in case.alternate_routes
            ],
        ],
    )


def _is_multi_code_case(case: XlsxEvalCase) -> bool:
    return len(_dedupe_codes(_source_expected_codes(case))) > 1


def _is_multi_code_adjusted(case: XlsxEvalCase) -> bool:
    source_codes = _dedupe_codes(_source_expected_codes(case))
    return source_codes != expected_codes_for_matching(source_codes)


def possible_multi_intent(case: XlsxEvalCase) -> bool:
    if not _is_multi_code_case(case):
        return False
    if not any(joiner in case.query for joiner in MULTI_INTENT_JOINERS):
        return False
    matched_terms = {
        term for term in MULTI_INTENT_RESOURCE_TERMS if term in case.query
    }
    return len(matched_terms) > 1


def min_code(codes: list[str]) -> str | None:
    if not codes:
        return None
    return min(codes, key=_code_sort_key)


def write_xlsx_result(
    output_path: Path,
    records: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = "results"
    headers = [
        "row_index",
        "eval_mode",
        "history_turn_count",
        "query",
        "expected_code",
        "alternate_codes",
        "expected_codes",
        "effective_expected_codes",
        "history_codes",
        "history_predicted_codes",
        "history_loop_counts",
        "predicted_status",
        "predicted_skill_id",
        "predicted_intent",
        "predicted_code",
        "matched",
        "match_reason",
        "strict_matched",
        "strict_match_reason",
        "source_expected_matched",
        "source_expected_match_reason",
        "possible_multi_intent",
        "elapsed_ms",
        "total_elapsed_ms",
        "loop_count",
        "correction_scopes",
        "resolved_query",
        "context_relation",
        "question",
        "options",
        "reason",
        "error_type",
        "error",
        "turn_results",
        "first_candidate_code",
        "final_candidate_code",
        "evaluator_verdicts",
        "fallback_used",
        "top_n_candidate_codes",
        "first_candidate_matched",
        "top_n_recalled",
        "full_route_matched",
        "evaluator_action",
        "expansion_rounds",
        "trace",
        "no_loop_predicted_status",
        "no_loop_predicted_skill_id",
        "no_loop_predicted_intent",
        "no_loop_predicted_code",
        "no_loop_matched",
        "no_loop_match_reason",
        "no_loop_elapsed_ms",
        "no_loop_total_elapsed_ms",
        "no_loop_loop_count",
        "no_loop_correction_scopes",
        "no_loop_resolved_query",
        "no_loop_context_relation",
        "no_loop_question",
        "no_loop_options",
        "no_loop_reason",
        "no_loop_error_type",
        "no_loop_error",
        "no_loop_turn_results",
        "no_loop_first_candidate_code",
        "no_loop_final_candidate_code",
        "no_loop_evaluator_verdicts",
        "no_loop_fallback_used",
        "no_loop_trace",
    ]
    result_sheet.append(headers)
    for record in records:
        result_sheet.append([_excel_cell(record.get(header)) for header in headers])

    summary_sheet = workbook.create_sheet("summary")
    summary_sheet.append(["metric", "value"])
    for key, value in summary.items():
        summary_sheet.append([key, _excel_cell(value)])

    workbook.save(output_path)


def normalize_code(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return str(value).strip()
        value = int(value)
    if isinstance(value, int):
        if value == 0:
            return ORDINARY_DIALOGUE_CODE
        if 0 < value < 1000:
            return _equivalent_code(f"{value:03d}")
        if 10000 <= value < 100000:
            return _equivalent_code(f"{value:06d}")
        return _equivalent_code(str(value))

    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"0+", text):
        return ORDINARY_DIALOGUE_CODE
    decimal_integer_match = re.fullmatch(r"(\d+)\.0+", text)
    if re.fullmatch(r"\d+", text) or decimal_integer_match:
        digits = decimal_integer_match.group(1) if decimal_integer_match else text
        number = int(digits)
        if number == 0:
            return ORDINARY_DIALOGUE_CODE
        if len(digits) < 3 and number < 1000:
            return _equivalent_code(f"{number:03d}")
        if len(digits) == 5 and 10000 <= number < 100000:
            return _equivalent_code(f"{number:06d}")
        return _equivalent_code(digits)
    return _equivalent_code(text)


def _equivalent_code(code: str) -> str:
    return LEGACY_CODE_EQUIVALENTS.get(code, code)


def _summary_for_expected_code(
    code: str,
    registry: SkillRegistry,
    code_index: dict[str, tuple[SkillDefinition, str]],
) -> DialogueRouteSummary:
    if code == ORDINARY_DIALOGUE_CODE:
        return DialogueRouteSummary(
            status="matched",
            intent="普通对话",
            code=ORDINARY_DIALOGUE_CODE,
        )

    if code == SEARCH_022_CODE and registry.has(SEARCH_SKILL_ID, active_only=True):
        skill = registry.get(SEARCH_SKILL_ID)
        return DialogueRouteSummary(
            status="matched",
            skill_id=skill.id,
            skill_name=skill.name,
            intent="搜索类历史期望",
            code=SEARCH_022_CODE,
        )

    indexed = code_index.get(code)
    if indexed is None:
        return DialogueRouteSummary(
            status="matched",
            intent=f"历史意图{code}",
            code=code,
        )

    skill, intent_name = indexed
    return DialogueRouteSummary(
        status="matched",
        skill_id=skill.id,
        skill_name=skill.name,
        intent=intent_name,
        code=code,
    )


def _build_code_index(
    registry: SkillRegistry,
) -> dict[str, tuple[SkillDefinition, str]]:
    code_index: dict[str, tuple[SkillDefinition, str]] = {}
    if registry.has(SEARCH_SKILL_ID, active_only=True):
        search_skill = registry.get(SEARCH_SKILL_ID)
        for intent_name, intent_schema in search_skill.intents.items():
            if (
                intent_schema.status == "active"
                and re.fullmatch(r"01\d", intent_schema.code)
            ):
                code_index.setdefault(intent_schema.code, (search_skill, intent_name))

    for card in registry.cards():
        skill = registry.get(card.id)
        for intent_name, intent_schema in skill.intents.items():
            if intent_schema.status == "active":
                code_index.setdefault(intent_schema.code, (skill, intent_name))
    return code_index


def _search_022_equivalent_codes(registry: SkillRegistry) -> set[str]:
    if not registry.has(SEARCH_SKILL_ID, active_only=True):
        return set()
    search_skill = registry.get(SEARCH_SKILL_ID)
    return {
        intent_schema.code
        for intent_schema in search_skill.intents.values()
        if intent_schema.status == "active" and re.fullmatch(r"01\d", intent_schema.code)
    }


def _build_record(
    *,
    case: XlsxEvalCase,
    result: RouteResult,
    matched: bool,
    match_reason: str,
    elapsed_ms: float,
    total_elapsed_ms: float,
    eval_mode: str,
    turn_results: list[dict[str, Any]],
    trace_summary: dict[str, Any],
    candidate_metrics: dict[str, Any] | None = None,
    strict_matched: bool = False,
    strict_match_reason: str | None = None,
    source_expected_matched: bool = False,
    source_expected_match_reason: str | None = None,
) -> dict[str, Any]:
    skill_id = result.skill.id if result.skill else None
    candidate_metrics = candidate_metrics or {}
    history_turn_results = [
        turn_result for turn_result in turn_results if not turn_result.get("is_scored")
    ]
    return {
        "row_index": case.row_index,
        "eval_mode": eval_mode,
        "history_turn_count": len(case.history_turns),
        "query": case.query,
        "expected_code": case.expected_code,
        "expected_skill_id": case.expected_skill_id,
        "expected_intent": case.expected_intent,
        "expected_params": case.expected_params,
        "alternate_routes": case.alternate_routes,
        "alternate_codes": case.alternate_codes,
        "expected_codes": case.source_expected_codes
        or [case.expected_code, *case.alternate_codes],
        "effective_expected_codes": [case.expected_code, *case.alternate_codes],
        "history_codes": [turn.expected_code for turn in case.history_turns],
        "history_predicted_codes": [
            turn_result.get("predicted_code") for turn_result in history_turn_results
        ],
        "history_loop_counts": [
            turn_result.get("loop_count") for turn_result in history_turn_results
        ],
        "predicted_status": result.status,
        "predicted_skill_id": skill_id,
        "predicted_intent": result.intent,
        "predicted_code": result.code,
        "predicted_params": result.params,
        "parameter_exact_matched": (
            result.params == case.expected_params
            if case.expected_params is not None
            else None
        ),
        "matched": matched,
        "match_reason": match_reason,
        "strict_matched": strict_matched,
        "strict_match_reason": strict_match_reason,
        "source_expected_matched": source_expected_matched,
        "source_expected_match_reason": source_expected_match_reason,
        "possible_multi_intent": possible_multi_intent(case),
        "elapsed_ms": round(elapsed_ms, 3),
        "total_elapsed_ms": round(total_elapsed_ms, 3),
        "loop_count": result.loop_count,
        "correction_scopes": result.correction_scopes,
        "resolved_query": result.resolved_query,
        "context_relation": result.context_relation,
        "question": result.question,
        "options": result.options,
        "reason": result.reason,
        "turn_results": turn_results,
        "first_candidate_code": trace_summary.get("first_candidate_code"),
        "final_candidate_code": trace_summary.get("final_candidate_code"),
        "evaluator_verdicts": trace_summary.get("evaluator_verdicts", []),
        "fallback_used": trace_summary.get("fallback_used", False),
        "top_n_candidate_codes": candidate_metrics.get("top_n_candidate_codes", []),
        "first_candidate_matched": candidate_metrics.get("first_candidate_matched"),
        "top_n_recalled": candidate_metrics.get("top_n_recalled"),
        "full_route_matched": candidate_metrics.get("full_route_matched"),
        "evaluator_action": (
            result.diagnostics.evaluator_action if result.diagnostics else None
        ),
        "expansion_rounds": (
            result.diagnostics.expansion_rounds if result.diagnostics else 0
        ),
    }


def _build_error_record(
    *,
    case: XlsxEvalCase,
    error: Exception,
    elapsed_ms: float,
    total_elapsed_ms: float,
    eval_mode: str,
) -> dict[str, Any]:
    return {
        "row_index": case.row_index,
        "eval_mode": eval_mode,
        "history_turn_count": len(case.history_turns),
        "query": case.query,
        "expected_code": case.expected_code,
        "expected_skill_id": case.expected_skill_id,
        "expected_intent": case.expected_intent,
        "expected_params": case.expected_params,
        "alternate_routes": case.alternate_routes,
        "alternate_codes": case.alternate_codes,
        "expected_codes": case.source_expected_codes
        or [case.expected_code, *case.alternate_codes],
        "effective_expected_codes": [case.expected_code, *case.alternate_codes],
        "history_codes": [turn.expected_code for turn in case.history_turns],
        "history_predicted_codes": [],
        "history_loop_counts": [],
        "predicted_status": "error",
        "predicted_skill_id": None,
        "predicted_intent": None,
        "predicted_code": None,
        "matched": False,
        "match_reason": "exception",
        "strict_matched": False,
        "strict_match_reason": "exception",
        "source_expected_matched": False,
        "source_expected_match_reason": "exception",
        "possible_multi_intent": possible_multi_intent(case),
        "elapsed_ms": round(elapsed_ms, 3),
        "total_elapsed_ms": round(total_elapsed_ms, 3),
        "loop_count": 0,
        "correction_scopes": [],
        "resolved_query": None,
        "context_relation": None,
        "question": None,
        "options": [],
        "reason": None,
        "error_type": type(error).__name__,
        "error": str(error),
        "turn_results": [],
        "first_candidate_code": None,
        "final_candidate_code": None,
        "evaluator_verdicts": [],
        "fallback_used": False,
    }


def _build_no_loop_record_fields(
    *,
    result: RouteResult,
    matched: bool,
    match_reason: str,
    elapsed_ms: float,
    total_elapsed_ms: float,
    turn_results: list[dict[str, Any]],
    trace: Any,
    trace_summary: dict[str, Any] | None = None,
) -> dict[str, Any]:
    skill_id = result.skill.id if result.skill else None
    trace_summary = trace_summary or {}
    return {
        "no_loop_predicted_status": result.status,
        "no_loop_predicted_skill_id": skill_id,
        "no_loop_predicted_intent": result.intent,
        "no_loop_predicted_code": result.code,
        "no_loop_matched": matched,
        "no_loop_match_reason": match_reason,
        "no_loop_elapsed_ms": round(elapsed_ms, 3),
        "no_loop_total_elapsed_ms": round(total_elapsed_ms, 3),
        "no_loop_loop_count": result.loop_count,
        "no_loop_correction_scopes": result.correction_scopes,
        "no_loop_resolved_query": result.resolved_query,
        "no_loop_context_relation": result.context_relation,
        "no_loop_question": result.question,
        "no_loop_options": result.options,
        "no_loop_reason": result.reason,
        "no_loop_error_type": None,
        "no_loop_error": None,
        "no_loop_turn_results": turn_results,
        "no_loop_first_candidate_code": trace_summary.get("first_candidate_code"),
        "no_loop_final_candidate_code": trace_summary.get("final_candidate_code"),
        "no_loop_evaluator_verdicts": trace_summary.get("evaluator_verdicts", []),
        "no_loop_fallback_used": trace_summary.get("fallback_used", False),
        "no_loop_trace": trace,
    }


def _build_no_loop_error_fields(
    *,
    error: Exception,
    elapsed_ms: float,
    total_elapsed_ms: float,
) -> dict[str, Any]:
    return {
        "no_loop_predicted_status": "error",
        "no_loop_predicted_skill_id": None,
        "no_loop_predicted_intent": None,
        "no_loop_predicted_code": None,
        "no_loop_matched": False,
        "no_loop_match_reason": "exception",
        "no_loop_elapsed_ms": round(elapsed_ms, 3),
        "no_loop_total_elapsed_ms": round(total_elapsed_ms, 3),
        "no_loop_loop_count": 0,
        "no_loop_correction_scopes": [],
        "no_loop_resolved_query": None,
        "no_loop_context_relation": None,
        "no_loop_question": None,
        "no_loop_options": [],
        "no_loop_reason": None,
        "no_loop_error_type": type(error).__name__,
        "no_loop_error": str(error),
        "no_loop_turn_results": [],
        "no_loop_first_candidate_code": None,
        "no_loop_final_candidate_code": None,
        "no_loop_evaluator_verdicts": [],
        "no_loop_fallback_used": False,
        "no_loop_trace": None,
    }


def _candidate_metrics(
    result: RouteResult,
    *,
    case: XlsxEvalCase,
    code_index: dict[str, tuple[SkillDefinition, str]],
    search_equivalent_codes: set[str],
) -> dict[str, Any]:
    candidates = _route_result_candidates(result)
    first_candidate = _candidate_by_id(
        candidates,
        result.diagnostics.first_candidate_id if result.diagnostics else None,
    )
    first_candidate = first_candidate or (candidates[0] if candidates else None)
    first_matched = _candidate_matches_expected_codes(
        first_candidate,
        case=case,
        search_equivalent_codes=search_equivalent_codes,
    )
    top_n_recalled = any(
        _candidate_matches_expected_codes(
            candidate,
            case=case,
            search_equivalent_codes=search_equivalent_codes,
        )
        for candidate in candidates
    )
    final_matched = result_matches_expected_codes(
        result,
        expected_code=case.expected_code,
        alternate_codes=_case_alternate_codes(case),
        search_equivalent_codes=search_equivalent_codes,
        allow_022_search_equivalence=case.allow_022_search_equivalence,
    )["matched"]
    full_route_matched = _full_route_matches_expected(
        result,
        case=case,
        code_index=code_index,
        code_matched=bool(final_matched),
    )
    evaluator_action = result.diagnostics.evaluator_action if result.diagnostics else None
    evaluator_called = bool(
        result.diagnostics and result.diagnostics.evaluator_called
    )
    expanded = bool(result.diagnostics and result.diagnostics.expansion_rounds > 0)
    expected_route = _expected_route(case, code_index)
    expected_skill_id = expected_route.get("skill_id") if expected_route else None
    expected_intent = expected_route.get("intent") if expected_route else None
    skill_hits = [
        candidate.get("skill_id") == expected_skill_id
        for candidate in candidates
    ] if expected_skill_id is not None else [False for _ in candidates]
    intent_hits = [
        candidate.get("skill_id") == expected_skill_id
        and candidate.get("intent") == expected_intent
        for candidate in candidates
    ] if expected_intent is not None else [False for _ in candidates]
    return {
        "top_n_candidate_codes": [candidate.get("code") for candidate in candidates],
        "first_candidate_matched": first_matched,
        "top_n_recalled": top_n_recalled,
        "full_route_matched": full_route_matched,
        "evaluator_switched": evaluator_action == "switch",
        "evaluator_switch_gain": (
            evaluator_action == "switch" and not first_matched and final_matched
        ),
        "evaluator_switch_loss": (
            evaluator_action == "switch" and first_matched and not final_matched
        ),
        "expanded": expanded,
        "expansion_gain": expanded and not first_matched and final_matched,
        "expansion_waste": expanded and not (not first_matched and final_matched),
        "evaluator_called": evaluator_called,
        "unnecessary_evaluator_call": (
            evaluator_called
            and evaluator_action == "select_top"
            and bool(first_matched)
        ),
        "skill_recall_at_1": any(skill_hits[:1]),
        "skill_recall_at_3": any(skill_hits[:3]),
        "skill_recall_at_5": any(skill_hits[:5]),
        "intent_recall_at_1": any(intent_hits[:1]),
        "intent_recall_at_2": any(intent_hits[:2]),
    }


def _expected_route(
    case: XlsxEvalCase,
    code_index: dict[str, tuple[SkillDefinition, str]],
) -> dict[str, str] | None:
    if case.expected_skill_id and case.expected_intent:
        return {
            "skill_id": case.expected_skill_id,
            "intent": case.expected_intent,
            "code": case.expected_code,
        }
    indexed = code_index.get(case.expected_code)
    if indexed is None:
        return None
    skill, intent = indexed
    return {"skill_id": skill.id, "intent": intent, "code": case.expected_code}


def _route_result_candidates(result: RouteResult) -> list[dict[str, Any]]:
    final_candidate = {
        "candidate_id": result.diagnostics.final_candidate_id
        if result.diagnostics
        else None,
        "skill_id": result.skill.id if result.skill else None,
        "intent": result.intent,
        "code": result.code,
        "params": result.params,
    }
    alternatives = [
        {
            "candidate_id": candidate.candidate_id,
            "skill_id": candidate.skill_id,
            "intent": candidate.intent,
            "code": candidate.code,
            "params": candidate.params,
        }
        for candidate in result.alternatives
    ]
    candidates = [final_candidate, *alternatives]
    if result.diagnostics and result.diagnostics.first_candidate_id:
        first = _candidate_by_id(candidates, result.diagnostics.first_candidate_id)
        if first is not None:
            candidates = [first, *[
                candidate
                for candidate in candidates
                if candidate.get("candidate_id") != first.get("candidate_id")
            ]]
    return candidates


def _candidate_by_id(
    candidates: list[dict[str, Any]],
    candidate_id: str | None,
) -> dict[str, Any] | None:
    if not candidate_id:
        return None
    for candidate in candidates:
        if candidate.get("candidate_id") == candidate_id:
            return candidate
    return None


def _candidate_matches_expected_codes(
    candidate: dict[str, Any] | None,
    *,
    case: XlsxEvalCase,
    search_equivalent_codes: set[str],
) -> bool:
    if candidate is None:
        return False
    result = RouteResult(
        status="matched",
        skill=(
            SkillRef(id=candidate["skill_id"], name=candidate["skill_id"])
            if candidate.get("skill_id")
            else None
        ),
        intent=candidate.get("intent"),
        code=candidate.get("code"),
        params=candidate.get("params") or {},
    )
    return bool(
        result_matches_expected_codes(
            result,
            expected_code=case.expected_code,
            alternate_codes=_case_alternate_codes(case),
            search_equivalent_codes=search_equivalent_codes,
            allow_022_search_equivalence=case.allow_022_search_equivalence,
        )["matched"],
    )


def _full_route_matches_expected(
    result: RouteResult,
    *,
    case: XlsxEvalCase,
    code_index: dict[str, tuple[SkillDefinition, str]],
    code_matched: bool,
) -> bool:
    explicit_routes = [
        {
            "skill_id": case.expected_skill_id,
            "intent": case.expected_intent,
            "code": case.expected_code,
        },
        *case.alternate_routes,
    ]
    explicit_routes = [
        route
        for route in explicit_routes
        if route.get("skill_id") and route.get("intent") and route.get("code")
    ]
    if explicit_routes:
        predicted_skill_id = result.skill.id if result.skill else None
        return any(
            predicted_skill_id == route["skill_id"]
            and result.intent == route["intent"]
            and result.code == route["code"]
            for route in explicit_routes
        )
    if not code_matched:
        return False
    accepted_codes = [case.expected_code, *case.alternate_codes]
    predicted_skill_id = result.skill.id if result.skill else None
    for code in accepted_codes:
        if code == ORDINARY_DIALOGUE_CODE:
            if predicted_skill_id is None and result.intent == "普通对话":
                return True
            continue
        indexed = code_index.get(code)
        if indexed is None:
            continue
        skill, intent_name = indexed
        if (
            predicted_skill_id == skill.id
            and result.intent == intent_name
            and result.code == code
        ):
            return True
    return code_matched and not any(code in code_index for code in accepted_codes)


def _has_strong_business_action(query: str) -> bool:
    return any(action in query for action in STRONG_BUSINESS_ACTIONS)


def _trace_summary(trace: list[dict[str, Any]]) -> dict[str, Any]:
    candidate_codes: list[str] = []
    evaluator_verdicts: list[dict[str, Any]] = []
    fallback_used = False
    final_candidate_code: str | None = None
    for event in trace:
        event_name = event.get("event")
        if event_name == "intent_candidates":
            for candidate in event.get("candidates") or []:
                code = candidate.get("code")
                if code:
                    candidate_codes.append(code)
        elif event_name == "first_candidate":
            candidate = event.get("candidate") or {}
            code = candidate.get("code")
            if code:
                candidate_codes.append(code)
        elif event_name == "rerank":
            evaluation = event.get("decision") or {}
            evaluator_verdicts.append(
                {
                    "verdict": evaluation.get("verdict"),
                    "selected_candidate_id": evaluation.get("selected_candidate_id"),
                    "expand_scope": evaluation.get("expand_scope"),
                    "confidence": evaluation.get("confidence"),
                },
            )
        elif event_name == "switch_gate":
            diagnostics = event.get("diagnostics") or {}
            fallback_used = diagnostics.get("evaluator_action") == "fallback"
        elif event_name == "result":
            result = event.get("result") or {}
            final_candidate_code = result.get("code")
        elif event_name in {"fallback"}:
            fallback_used = True

    return {
        "first_candidate_code": candidate_codes[0] if candidate_codes else None,
        "final_candidate_code": final_candidate_code,
        "evaluator_verdicts": evaluator_verdicts,
        "fallback_used": fallback_used,
    }


def _build_turn_result(
    *,
    turn_index: int,
    query: str,
    expected_code: str,
    result: RouteResult,
    elapsed_ms: float,
    is_scored: bool,
) -> dict[str, Any]:
    skill_id = result.skill.id if result.skill else None
    return {
        "turn_index": turn_index,
        "query": query,
        "expected_code": expected_code,
        "predicted_status": result.status,
        "predicted_skill_id": skill_id,
        "predicted_intent": result.intent,
        "predicted_code": result.code,
        "loop_count": result.loop_count,
        "correction_scopes": result.correction_scopes,
        "elapsed_ms": round(elapsed_ms, 3),
        "resolved_query": result.resolved_query,
        "context_relation": result.context_relation,
        "question": result.question,
        "options": result.options,
        "reason": result.reason,
        "is_scored": is_scored,
    }


def _history_column_pairs(
    headers: list[str],
    header_to_index: dict[str, int],
) -> list[tuple[int, int]]:
    pairs = []
    for index, header in enumerate(headers):
        match = re.fullmatch(r"第(.+)轮对话", header)
        if not match:
            continue
        expected_header = f"第{match.group(1)}轮预期意图"
        expected_index = header_to_index.get(expected_header)
        if expected_index is None:
            continue
        pairs.append((index, expected_index))
    return pairs


def _format_history_column_pairs(
    headers: list[str],
    header_to_index: dict[str, int],
) -> list[tuple[int, int]]:
    pairs: list[tuple[int, int, int]] = []
    for index, header in enumerate(headers):
        match = re.fullmatch(r"上文(\d+)", header)
        if not match:
            continue
        turn_number = int(match.group(1))
        expected_index = header_to_index.get(f"上文意图{turn_number}")
        if expected_index is None:
            continue
        pairs.append((turn_number, index, expected_index))
    return [
        (query_index, expected_index)
        for _, query_index, expected_index in sorted(pairs, key=lambda item: item[0])
    ]


def _require_header(header_to_index: dict[str, int], header: str) -> int:
    index = header_to_index.get(header)
    if index is None:
        raise ValueError(f"Excel 缺少必需表头: {header}")
    return index


def _require_any_header(
    header_to_index: dict[str, int],
    headers: tuple[str, ...],
) -> int:
    for header in headers:
        index = header_to_index.get(header)
        if index is not None:
            return index
    raise ValueError(f"Excel 缺少必需表头之一: {', '.join(headers)}")


def _first_header_index(
    header_to_index: dict[str, int],
    headers: tuple[str, ...],
) -> int | None:
    return next(
        (
            header_to_index[header]
            for header in headers
            if header in header_to_index
        ),
        None,
    )


def _optional_cell(row: tuple[Any, ...], index: int | None) -> str | None:
    if index is None:
        return None
    value = _cell_text(_row_value(row, index))
    return value or None


def _parse_expected_params(value: Any) -> dict[str, Any] | None:
    text = _cell_text(value)
    if not text:
        return None
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"expected_params 必须是 JSON object: {text}") from exc
    if not isinstance(parsed, dict):
        raise ValueError("expected_params 必须是 JSON object")
    return parsed


def _parse_alternate_routes(value: Any) -> list[dict[str, str]]:
    text = _cell_text(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"alternate_routes 必须是 JSON array: {text}") from exc
    if not isinstance(parsed, list) or not all(isinstance(item, dict) for item in parsed):
        raise ValueError("alternate_routes 必须是 JSON object array")
    routes: list[dict[str, str]] = []
    for item in parsed:
        route = {
            key: str(item[key])
            for key in ("skill_id", "intent", "code")
            if item.get(key) is not None
        }
        if {"skill_id", "intent", "code"} <= route.keys():
            routes.append(route)
    return routes


def _row_value(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _average(values: list[float] | list[int]) -> float:
    if not values:
        return 0.0
    return sum(values) / len(values)


def _percentile(values: list[float], percentile: int) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    position = (len(ordered) - 1) * percentile / 100
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    fraction = position - lower
    return ordered[lower] + (ordered[upper] - ordered[lower]) * fraction


def _dedupe_codes(codes: list[str | None]) -> list[str]:
    deduped = []
    for code in codes:
        if code and code not in deduped:
            deduped.append(code)
    return deduped


def _code_sort_key(code: str) -> tuple[int, int | str]:
    if re.fullmatch(r"\d+", code):
        return (0, int(code))
    return (1, code)


def _excel_cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        import json

        return json.dumps(value, ensure_ascii=False)
    return value
